"""Source parsers for text, markdown, CSV, XLSX, SQLite, and JSON/JSONL."""

from __future__ import annotations

import csv
import json
import sqlite3
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Union
from luna.ingest.chunker import chunk_text
from luna.ingest.models import ParsedChunk


@dataclass
class TableData:
    """Parsed in-memory tabular dataset."""
    name: str
    source_id: str
    columns: list[str]
    rows: list[dict[str, Any]]
    total_rows: int
    foreign_keys: list[dict[str, Any]] = field(default_factory=list)
    truncated: bool = False


def flatten_json_object(obj: dict[str, Any], prefix: str = "", depth: int = 1, max_depth: int = 3) -> dict[str, Any]:
    """Flatten nested dictionary up to max_depth using dotted keys."""
    items: dict[str, Any] = {}
    for k, v in obj.items():
        key = f"{prefix}{k}" if not prefix else f"{prefix}.{k}"
        if isinstance(v, dict) and depth < max_depth:
            items.update(flatten_json_object(v, prefix=key, depth=depth + 1, max_depth=max_depth))
        elif isinstance(v, list):
            items[key] = json.dumps(v)
        else:
            items[key] = v
    return items


def parse_source_file(
    source_id: str,
    name: str,
    kind: str,
    file_path: Union[str, Path],
    max_table_rows: int = 200_000,
) -> tuple[list[ParsedChunk], list[TableData], list[str]]:
    """Parse a file based on its kind.

    Returns:
        (chunks, tables, warnings)
    """
    path = Path(file_path)
    warnings: list[str] = []
    chunks: list[ParsedChunk] = []
    tables: list[TableData] = []

    if kind in ("text", "txt", "md"):
        try:
            raw_text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            raw_text = path.read_text(encoding="latin-1", errors="replace")
        chunks = chunk_text(source_id, name, raw_text)
        return chunks, tables, warnings

    elif kind == "csv":
        try:
            content = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            content = path.read_text(encoding="latin-1", errors="replace")

        # Delimiter sniffing
        sample = content[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample)
            delimiter = dialect.delimiter
        except Exception:
            delimiter = ","

        reader = csv.reader(content.splitlines(), delimiter=delimiter)
        rows_raw = list(reader)
        if not rows_raw:
            return chunks, tables, warnings

        header = [c.strip() for c in rows_raw[0] if c.strip()]
        if not header:
            raise ValueError(f"CSV file '{name}' must have a non-empty header row")

        data_rows: list[dict[str, Any]] = []
        truncated = False
        for idx, row in enumerate(rows_raw[1:], start=1):
            if idx > max_table_rows:
                truncated = True
                warnings.append(f"Table '{name}' was truncated at {max_table_rows:,} rows")
                break
            record = {}
            for col_idx, col_name in enumerate(header):
                val = row[col_idx] if col_idx < len(row) else ""
                record[col_name] = val.strip()
            data_rows.append(record)

        table_name = Path(name).stem
        tables.append(
            TableData(
                name=table_name,
                source_id=source_id,
                columns=header,
                rows=data_rows,
                total_rows=len(rows_raw) - 1,
                truncated=truncated,
            )
        )
        return chunks, tables, warnings

    elif kind in ("xlsx", "xls"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)
            header = None
            data_rows = []
            truncated = False

            row_count = 0
            for r in rows_iter:
                # Find first non-empty row as header
                if header is None:
                    if any(cell is not None for cell in r):
                        header = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(r)]
                    continue

                if not any(cell is not None for cell in r):
                    continue

                row_count += 1
                if row_count > max_table_rows:
                    truncated = True
                    warnings.append(f"Sheet '{sheet_name}' was truncated at {max_table_rows:,} rows")
                    break

                row_dict = {}
                for col_idx, col_name in enumerate(header):
                    val = r[col_idx] if col_idx < len(r) else None
                    row_dict[col_name] = str(val).strip() if val is not None else ""
                data_rows.append(row_dict)

            if header and data_rows:
                tbl_name = f"{Path(name).stem}_{sheet_name}" if len(wb.sheetnames) > 1 else Path(name).stem
                tables.append(
                    TableData(
                        name=tbl_name,
                        source_id=source_id,
                        columns=header,
                        rows=data_rows,
                        total_rows=row_count,
                        truncated=truncated,
                    )
                )
        return chunks, tables, warnings

    elif kind in ("sqlite", "db"):
        # Open in read-only URI mode
        uri = f"file:{path.resolve().as_posix()}?mode=ro"
        conn = sqlite3.connect(uri, uri=True)
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()

        # Find user tables
        cur.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")
        user_tables = [row[0] for row in cur.fetchall()]

        for tbl in user_tables:
            # Foreign keys
            cur.execute(f"PRAGMA foreign_key_list('{tbl}')")
            fk_rows = cur.fetchall()
            foreign_keys = [
                {"table": r[2], "from": r[3], "to": r[4]}
                for r in fk_rows
            ]

            # Columns
            cur.execute(f"PRAGMA table_info('{tbl}')")
            columns = [r[1] for r in cur.fetchall()]

            # Fetch rows with LIMIT
            cur.execute(f"SELECT COUNT(*) FROM '{tbl}'")
            total_cnt = cur.fetchone()[0]

            cur.execute(f"SELECT * FROM '{tbl}' LIMIT {max_table_rows}")
            rows = [dict(r) for r in cur.fetchall()]
            truncated = total_cnt > max_table_rows
            if truncated:
                warnings.append(f"SQLite table '{tbl}' was truncated at {max_table_rows:,} rows")

            tables.append(
                TableData(
                    name=tbl,
                    source_id=source_id,
                    columns=columns,
                    rows=rows,
                    total_rows=total_cnt,
                    foreign_keys=foreign_keys,
                    truncated=truncated,
                )
            )
        conn.close()
        return chunks, tables, warnings

    elif kind == "jsonl":
        data_rows = []
        truncated = False
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            for idx, line in enumerate(f, start=1):
                line_str = line.strip()
                if not line_str:
                    continue
                if idx > max_table_rows:
                    truncated = True
                    warnings.append(f"JSONL '{name}' truncated at {max_table_rows:,} rows")
                    break
                try:
                    obj = json.loads(line_str)
                    if isinstance(obj, dict):
                        data_rows.append(flatten_json_object(obj))
                except Exception:
                    pass

        if data_rows:
            all_cols: list[str] = []
            for d in data_rows:
                for k in d.keys():
                    if k not in all_cols:
                        all_cols.append(k)
            tables.append(
                TableData(
                    name=Path(name).stem,
                    source_id=source_id,
                    columns=all_cols,
                    rows=data_rows,
                    total_rows=len(data_rows),
                    truncated=truncated,
                )
            )
        return chunks, tables, warnings

    elif kind == "json":
        raw_text = path.read_text(encoding="utf-8", errors="replace")
        try:
            data = json.loads(raw_text)
        except Exception:
            # Fall back to text chunking
            chunks = chunk_text(source_id, name, raw_text)
            return chunks, tables, warnings

        if isinstance(data, list) and all(isinstance(x, dict) for x in data[:10]):
            # Top-level list of objects -> single table
            flattened = [flatten_json_object(x) for x in data[:max_table_rows]]
            cols = []
            for d in flattened:
                for k in d.keys():
                    if k not in cols:
                        cols.append(k)
            tables.append(
                TableData(
                    name=Path(name).stem,
                    source_id=source_id,
                    columns=cols,
                    rows=flattened,
                    total_rows=len(data),
                    truncated=len(data) > max_table_rows,
                )
            )
        elif isinstance(data, dict) and any(isinstance(v, list) and v and isinstance(v[0], dict) for v in data.values()):
            # Multi-table dictionary
            for k, val in data.items():
                if isinstance(val, list) and val and isinstance(val[0], dict):
                    flattened = [flatten_json_object(x) for x in val[:max_table_rows]]
                    cols = []
                    for d in flattened:
                        for col in d.keys():
                            if col not in cols:
                                cols.append(col)
                    tables.append(
                        TableData(
                            name=f"{Path(name).stem}_{k}",
                            source_id=source_id,
                            columns=cols,
                            rows=flattened,
                            total_rows=len(val),
                            truncated=len(val) > max_table_rows,
                        )
                    )
        else:
            # Pretty-printed text chunking
            pretty = json.dumps(data, indent=2)
            chunks = chunk_text(source_id, name, pretty)

        return chunks, tables, warnings

    else:
        # Default fallback to text chunking
        try:
            raw_text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            raw_text = path.read_text(encoding="latin-1", errors="replace")
        chunks = chunk_text(source_id, name, raw_text)
        return chunks, tables, warnings
